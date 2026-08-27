"""
Extração de texto de múltiplos formatos do corpus Ocultismo.
Suporta: PDF, DOCX, ODT, TXT, RTF, JSON, imagens (OCR via easyocr).
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from odf.opendocument import load as odf_load
from odf.text import P as OdfP
from striprtf.striprtf import rtf_to_text

logger = logging.getLogger(__name__)

SOURCE_ROOT = Path("/mnt/e/Ocultismo")

SKIP_IMAGE_DIRS = {"Pontos"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
TEXT_EXTS = {".pdf", ".docx", ".odt", ".txt", ".rtf", ".doc", ".json"}

CATEGORY_MAP = {
    "Angelical e olimpicos": "magia_angelical",
    "Apometria": "apometria",
    "Backup": "backup",
    "Fitoterapia": "fitoterapia",
    "Judaismo": "judaismo",
    "Kundalini": "kundalini",
    "Livros": "livros",
    "Magia Angelical e Olimpicos": "magia_angelical",
    "Magia Cigana": "magia_cigana",
    "Magia da 7 Cruzes": "magia_7_cruzes",
    "Magia das Ervas": "magia_ervas",
    "Magia do Fogo": "magia_fogo",
    "Oracoes e Rezas": "oracoes_rezas",
    "Oração Pinheiro": "oracao_pinheiro",
    "Pontos": "pontos_cantados",
    "Reiki": "reiki",
    "Sepher Ra Hazim": "sepher_ra_hazim",
    "Templo da Pena": "templo_pena",
    "resumos_individuais": "resumos",
}


@dataclass
class ExtractedDocument:
    doc_id: str
    title: str
    category: str
    subcategory: str
    source_path: str
    file_type: str
    text: str
    page_count: int = 0
    char_count: int = 0
    metadata: dict = field(default_factory=dict)


def compute_doc_id(path: Path) -> str:
    return hashlib.sha256(str(path).encode()).hexdigest()[:16]


def get_category(path: Path) -> tuple[str, str]:
    parts = path.relative_to(SOURCE_ROOT).parts
    top_dir = parts[0] if parts else ""
    category = CATEGORY_MAP.get(top_dir, "outros")
    subcategory = parts[1] if len(parts) > 2 else ""
    return category, subcategory


def extract_pdf(path: Path) -> tuple[str, int]:
    try:
        doc = pymupdf.open(str(path))
        pages = []
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                pages.append(text)
        doc.close()
        return "\n\n".join(pages), len(pages)
    except Exception as e:
        logger.error("PDF extraction failed for %s: %s", path.name, e)
        return "", 0


def extract_docx(path: Path) -> tuple[str, int]:
    try:
        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs), len(paragraphs)
    except Exception as e:
        logger.error("DOCX extraction failed for %s: %s", path.name, e)
        return "", 0


def extract_odt(path: Path) -> tuple[str, int]:
    try:
        doc = odf_load(str(path))
        paragraphs = [
            str(p) for p in doc.getElementsByType(OdfP) if str(p).strip()
        ]
        return "\n\n".join(paragraphs), len(paragraphs)
    except Exception as e:
        logger.error("ODT extraction failed for %s: %s", path.name, e)
        return "", 0


def extract_txt(path: Path) -> tuple[str, int]:
    try:
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                text = path.read_text(encoding=encoding)
                return text, text.count("\n") + 1
            except (UnicodeDecodeError, UnicodeError):
                continue
        return "", 0
    except Exception as e:
        logger.error("TXT extraction failed for %s: %s", path.name, e)
        return "", 0


def extract_rtf(path: Path) -> tuple[str, int]:
    try:
        raw = path.read_text(encoding="latin-1")
        text = rtf_to_text(raw)
        return text, text.count("\n") + 1
    except Exception as e:
        logger.error("RTF extraction failed for %s: %s", path.name, e)
        return "", 0


def extract_json(path: Path) -> tuple[str, int]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        text = json.dumps(data, ensure_ascii=False, indent=2)
        return text, 1
    except Exception as e:
        logger.error("JSON extraction failed for %s: %s", path.name, e)
        return "", 0


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".odt": extract_odt,
    ".txt": extract_txt,
    ".rtf": extract_rtf,
    ".json": extract_json,
}


def extract_text_file(path: Path) -> ExtractedDocument | None:
    ext = path.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        if ext == ".doc":
            logger.warning("Skipping .doc (legacy format): %s", path.name)
        return None

    text, page_count = extractor(path)
    if not text or len(text.strip()) < 50:
        logger.warning("Insufficient text extracted from %s (%d chars)", path.name, len(text.strip()))
        return None

    category, subcategory = get_category(path)
    title = path.stem
    # Clean long Anna's Archive metadata from filenames
    title = re.sub(r"\s*--.*$", "", title)

    return ExtractedDocument(
        doc_id=compute_doc_id(path),
        title=title,
        category=category,
        subcategory=subcategory,
        source_path=str(path),
        file_type=ext.lstrip("."),
        text=text,
        page_count=page_count,
        char_count=len(text),
    )


def extract_image_ocr(path: Path, ocr_reader) -> ExtractedDocument | None:
    try:
        results = ocr_reader.readtext(str(path), detail=0, paragraph=True)
        text = "\n".join(results)
        if not text or len(text.strip()) < 30:
            return None

        category, subcategory = get_category(path)
        return ExtractedDocument(
            doc_id=compute_doc_id(path),
            title=path.stem,
            category=category,
            subcategory=subcategory,
            source_path=str(path),
            file_type="image_ocr",
            text=text,
            page_count=1,
            char_count=len(text),
        )
    except Exception as e:
        logger.error("OCR failed for %s: %s", path.name, e)
        return None


def should_process_image(path: Path) -> bool:
    for skip_dir in SKIP_IMAGE_DIRS:
        if skip_dir in path.parts:
            return False
    return True


def collect_files(root: Path) -> tuple[list[Path], list[Path]]:
    text_files = []
    image_files = []

    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        ext = p.suffix.lower()
        if ext in TEXT_EXTS:
            text_files.append(p)
        elif ext in IMAGE_EXTS and should_process_image(p):
            image_files.append(p)

    return text_files, image_files


def run_extraction(root: Path = SOURCE_ROOT) -> list[ExtractedDocument]:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        stream=sys.stderr,
    )

    text_files, image_files = collect_files(root)
    logger.info("Found %d text files and %d images to process", len(text_files), len(image_files))

    documents: list[ExtractedDocument] = []

    # --- Phase 1: text files ---
    for i, path in enumerate(text_files, 1):
        logger.info("[%d/%d] Extracting: %s", i, len(text_files), path.name)
        doc = extract_text_file(path)
        if doc:
            documents.append(doc)

    logger.info("Extracted %d text documents", len(documents))

    # --- Phase 2: OCR images ---
    if image_files:
        logger.info("Loading EasyOCR reader (pt + en)...")
        import easyocr
        reader = easyocr.Reader(["pt", "en"], gpu=True)

        for i, path in enumerate(image_files, 1):
            logger.info("[%d/%d] OCR: %s", i, len(image_files), path.name)
            doc = extract_image_ocr(path, reader)
            if doc:
                documents.append(doc)

        logger.info("Total documents after OCR: %d", len(documents))

    # --- Summary ---
    total_chars = sum(d.char_count for d in documents)
    logger.info(
        "Extraction complete: %d documents, %.1f MB of text",
        len(documents),
        total_chars / 1_048_576,
    )
    return documents
